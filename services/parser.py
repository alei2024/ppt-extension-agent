"""
PPT解析服务
使用python-pptx进行PPT文件解析，提取文本、图片和层级结构
"""

from typing import Dict, List, Optional
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
import logging
import json

logger = logging.getLogger(__name__)


class PPTParser:
    """
    PPT文件解析器
    支持解析PPT的层级结构、文本内容、图片描述等
    """
    
    def __init__(self, image_output_dir: str = "./uploads/images", recognize_images: bool = True):
        """
        初始化解析器
        
        Args:
            image_output_dir: 图片保存目录
            recognize_images: 是否识别图片内容
        """
        self.image_output_dir = Path(image_output_dir)
        self.image_output_dir.mkdir(parents=True, exist_ok=True)
        self.recognize_images = recognize_images
        
        # 延迟导入图片识别服务（避免OCR库未安装时出错）
        self.image_recognizer = None
        if recognize_images:
            try:
                from services.image_recognizer import image_recognizer
                self.image_recognizer = image_recognizer
                logger.info("图片识别功能已启用")
            except Exception as e:
                logger.warning(f"图片识别功能不可用: {str(e)}")
                self.recognize_images = False
        
        logger.info("PPT解析器初始化完成")
    
    def parse_from_file(self, file_path: str) -> Dict:
        """
        从本地文件解析PPT
        
        Args:
            file_path: PPT文件路径
            
        Returns:
            解析后的数据结构，包含：
            - slides: 页面列表
            - structure: 层级结构
            - metadata: 元数据
        """
        try:
            logger.info(f"开始解析PPT文件: {file_path}")
            
            if not Path(file_path).exists():
                raise FileNotFoundError(f"PPT文件不存在: {file_path}")
            
            prs = Presentation(file_path)
            
            # 为每个PPT文件创建独立的图片目录
            ppt_name = Path(file_path).stem
            ppt_image_dir = self.image_output_dir / ppt_name
            ppt_image_dir.mkdir(parents=True, exist_ok=True)
            
            result = {
                "metadata": self._extract_metadata(prs, file_path),
                "slides": [],
                "structure": {
                    "title": "",
                    "sections": []
                }
            }
            
            for idx, slide in enumerate(prs.slides):
                slide_data = self._parse_slide(slide, idx + 1, ppt_image_dir)
                result["slides"].append(slide_data)
                
                if idx == 0:
                    result["structure"]["title"] = slide_data.get("title", "")
            
            logger.info(f"PPT解析完成，共{len(result['slides'])}页")
            return result
            
        except Exception as e:
            logger.error(f"PPT解析失败: {str(e)}")
            raise
    
    def parse_from_url(self, url: str, download_path: str = "./uploads") -> Dict:
        """
        从URL下载并解析PPT
        
        Args:
            url: PPT文件的URL地址
            download_path: 下载保存路径
            
        Returns:
            解析后的数据结构
        """
        try:
            import requests
            from urllib.parse import urlparse
            import os
            
            logger.info(f"开始从URL下载PPT: {url}")
            
            response = requests.get(url, stream=True, timeout=30)
            response.raise_for_status()
            
            filename = os.path.basename(urlparse(url).path)
            if not filename or not filename.endswith(('.pptx', '.ppt')):
                filename = f"downloaded_{hash(url)}.pptx"
            
            file_path = os.path.join(download_path, filename)
            
            os.makedirs(download_path, exist_ok=True)
            
            with open(file_path, 'wb') as f:
                for chunk in response.iter_content(chunk_size=8192):
                    f.write(chunk)
            
            logger.info(f"PPT下载完成: {file_path}")
            
            return self.parse_from_file(file_path)
            
        except Exception as e:
            logger.error(f"从URL解析PPT失败: {str(e)}")
            raise
    
    def _parse_slide(self, slide, page_number: int, image_dir: Path = None) -> Dict:
        """
        解析单个幻灯片
        
        Args:
            slide: python-pptx的Slide对象
            page_number: 页码
            image_dir: 图片保存目录
            
        Returns:
            幻灯片数据字典
        """
        slide_data = {
            "page_number": page_number,
            "title": "",
            "text_boxes": [],
            "images": [],
            "tables": [],
            "notes": ""
        }
        
        for shape in slide.shapes:
            logger.debug(f"处理形状: type={shape.shape_type}, id={shape.shape_id}")
            
            if shape.shape_type == 14:
                slide_data["title"] = shape.text.strip()
            
            elif shape.has_text_frame:
                text_box = self._parse_text_box(shape)
                if text_box["text"].strip():
                    slide_data["text_boxes"].append(text_box)
            
            elif shape.shape_type == 13:
                logger.info(f"发现图片: shape_id={shape.shape_id}, page={page_number}")
                image_data = self._parse_image(shape, page_number, image_dir)
                if image_data.get("file_path"):
                    logger.info(f"图片提取成功: {image_data['file_path']}")
                slide_data["images"].append(image_data)
            
            elif shape.has_table:
                table_data = self._parse_table(shape)
                slide_data["tables"].append(table_data)
        
        logger.info(f"页面 {page_number} 解析完成: {len(slide_data['text_boxes'])} 个文本框, {len(slide_data['images'])} 张图片")
        
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            slide_data["notes"] = slide.notes_slide.notes_text_frame.text.strip()
        
        return slide_data
    
    def _parse_text_box(self, shape) -> Dict:
        """
        解析文本框
        
        Args:
            shape: python-pptx的Shape对象
            
        Returns:
            文本框数据字典
        """
        text_frame = shape.text_frame
        paragraphs = []
        
        for paragraph in text_frame.paragraphs:
            para_data = {
                "text": paragraph.text.strip(),
                "level": paragraph.level,
                "font_size": None,
                "font_name": None,
                "bold": False,
                "italic": False
            }
            
            if paragraph.runs:
                run = paragraph.runs[0]
                if run.font.size:
                    para_data["font_size"] = run.font.size.pt
                if run.font.name:
                    para_data["font_name"] = run.font.name
                para_data["bold"] = run.font.bold
                para_data["italic"] = run.font.italic
            
            if para_data["text"]:
                paragraphs.append(para_data)
        
        return {
            "id": shape.shape_id,
            "text": text_frame.text.strip(),
            "paragraphs": paragraphs,
            "position": {
                "left": float(shape.left / Inches(1)),
                "top": float(shape.top / Inches(1)),
                "width": float(shape.width / Inches(1)),
                "height": float(shape.height / Inches(1))
            },
            "is_title": self._is_title_box(shape)
        }
    
    def _is_title_box(self, shape) -> bool:
        """
        判断是否为标题文本框
        
        Args:
            shape: python-pptx的Shape对象
            
        Returns:
            是否为标题
        """
        if not shape.has_text_frame:
            return False
        
        text_frame = shape.text_frame
        
        if text_frame.paragraphs and text_frame.paragraphs[0].runs:
            run = text_frame.paragraphs[0].runs[0]
            if run.font.size:
                font_size = run.font.size.pt
                return font_size > 32
        
        return False
    
    def _parse_image(self, shape, page_number: int, image_dir: Path = None) -> Dict:
        """
        解析图片并保存到本地，可选识别图片内容
        
        Args:
            shape: python-pptx的Shape对象
            page_number: 页码
            image_dir: 图片保存目录
            
        Returns:
            图片数据字典，包含：
            - id: 图片ID
            - file_path: 图片文件路径（相对路径）
            - position: 位置信息
            - description: 图片描述（识别结果）
            - ocr_text: OCR识别的文字
        """
        image_data = {
            "id": shape.shape_id,
            "position": {
                "left": float(shape.left / Inches(1)),
                "top": float(shape.top / Inches(1)),
                "width": float(shape.width / Inches(1)),
                "height": float(shape.height / Inches(1))
            },
            "file_path": "",
            "description": "",
            "ocr_text": ""
        }
        
        # 提取图片二进制数据并保存
        saved_file_path = None
        if image_dir and hasattr(shape, "image"):
            try:
                image = shape.image
                image_bytes = image.blob
                
                # 确定图片格式
                ext = image.ext
                if not ext:
                    ext = "png"  # 默认使用png
                
                # 生成文件名
                filename = f"slide_{page_number}_img_{shape.shape_id}.{ext}"
                file_path = image_dir / filename
                
                # 保存图片
                with open(file_path, "wb") as f:
                    f.write(image_bytes)
                
                saved_file_path = str(file_path)
                
                # 保存相对路径（相对于uploads目录）
                relative_path = f"images/{image_dir.name}/{filename}"
                image_data["file_path"] = relative_path
                
                logger.info(f"图片已保存: {file_path}")
                
            except Exception as e:
                logger.warning(f"提取图片失败 (shape_id={shape.shape_id}): {str(e)}")
        
        # 识别图片内容
        if self.recognize_images and self.image_recognizer and saved_file_path:
            try:
                recognition_result = self.image_recognizer.recognize_image(saved_file_path, page_number)
                image_data["description"] = recognition_result.get("description", "")
                image_data["ocr_text"] = recognition_result.get("text", "")
                logger.info(f"图片识别完成: {saved_file_path}")
            except Exception as e:
                logger.warning(f"图片识别失败: {str(e)}")
        
        return image_data
    
    def _parse_table(self, shape) -> Dict:
        """
        解析表格
        
        Args:
            shape: python-pptx的Shape对象
            
        Returns:
            表格数据字典
        """
        table = shape.table
        rows = []
        
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                row_data.append(cell.text.strip())
            rows.append(row_data)
        
        return {
            "id": shape.shape_id,
            "rows": len(rows),
            "columns": len(rows[0]) if rows else 0,
            "data": rows,
            "position": {
                "left": float(shape.left / Inches(1)),
                "top": float(shape.top / Inches(1)),
                "width": float(shape.width / Inches(1)),
                "height": float(shape.height / Inches(1))
            }
        }
    
    def _extract_metadata(self, prs, file_path: str) -> Dict:
        """
        提取PPT元数据
        
        Args:
            prs: python-pptx的Presentation对象
            file_path: 文件路径
            
        Returns:
            元数据字典
        """
        return {
            "file_name": Path(file_path).name,
            "file_path": file_path,
            "slide_count": len(prs.slides),
            "author": prs.core_properties.author or "",
            "title": prs.core_properties.title or "",
            "created": prs.core_properties.created.isoformat() if prs.core_properties.created else "",
            "modified": prs.core_properties.modified.isoformat() if prs.core_properties.modified else ""
        }
    
    def extract_structure(self, ppt_data: Dict) -> Dict:
        """
        提取PPT的层级结构
        
        Args:
            ppt_data: 原始PPT数据
            
        Returns:
            结构化的层级信息：
            - title: 标题
            - sections: 章节列表
            - subsections: 子章节
        """
        structure = {
            "title": ppt_data.get("structure", {}).get("title", ""),
            "sections": []
        }
        
        current_section = None
        
        for slide in ppt_data.get("slides", []):
            title = slide.get("title", "")
            
            if self._is_section_title(title):
                if current_section:
                    structure["sections"].append(current_section)
                current_section = {
                    "title": title,
                    "page_number": slide["page_number"],
                    "subsections": []
                }
            elif current_section:
                current_section["subsections"].append({
                    "title": title,
                    "page_number": slide["page_number"],
                    "content": [tb["text"] for tb in slide.get("text_boxes", [])]
                })
        
        if current_section:
            structure["sections"].append(current_section)
        
        return structure
    
    def _is_section_title(self, title: str) -> bool:
        """
        判断是否为章节标题
        
        Args:
            title: 标题文本
            
        Returns:
            是否为章节标题
        """
        section_keywords = ["章", "Chapter", "Part", "部分", "单元"]
        return any(keyword in title for keyword in section_keywords)
    
    def extract_images(self, ppt_data: Dict) -> List[Dict]:
        """
        提取PPT中的图片信息
        
        Args:
            ppt_data: 原始PPT数据
            
        Returns:
            图片信息列表，包含：
            - image_id: 图片ID
            - page_number: 所在页码
            - description: 图片描述
            - position: 位置信息
        """
        images = []
        
        for slide in ppt_data.get("slides", []):
            for image in slide.get("images", []):
                images.append({
                    "image_id": f"slide_{slide['page_number']}_img_{image['id']}",
                    "page_number": slide["page_number"],
                    "description": image.get("description", ""),
                    "position": image.get("position", {})
                })
        
        return images
    
    def extract_text_chunks(self, ppt_data: Dict, chunk_size: int = 500) -> List[Dict]:
        """
        提取文本切片（用于向量化）
        
        Args:
            ppt_data: 原始PPT数据
            chunk_size: 每个切片的最大字符数
            
        Returns:
            文本切片列表
        """
        chunks = []
        
        for slide in ppt_data.get("slides", []):
            slide_text = " ".join([tb["text"] for tb in slide.get("text_boxes", [])])
            
            if len(slide_text) <= chunk_size:
                chunks.append({
                    "content": slide_text,
                    "page_number": slide["page_number"],
                    "title": slide.get("title", ""),
                    "chunk_id": f"slide_{slide['page_number']}_chunk_0"
                })
            else:
                words = slide_text.split()
                current_chunk = ""
                chunk_index = 0
                
                for word in words:
                    if len(current_chunk) + len(word) + 1 <= chunk_size:
                        current_chunk += " " + word if current_chunk else word
                    else:
                        chunks.append({
                            "content": current_chunk,
                            "page_number": slide["page_number"],
                            "title": slide.get("title", ""),
                            "chunk_id": f"slide_{slide['page_number']}_chunk_{chunk_index}"
                        })
                        current_chunk = word
                        chunk_index += 1
                
                if current_chunk:
                    chunks.append({
                        "content": current_chunk,
                        "page_number": slide["page_number"],
                        "title": slide.get("title", ""),
                        "chunk_id": f"slide_{slide['page_number']}_chunk_{chunk_index}"
                    })
        
        return chunks
